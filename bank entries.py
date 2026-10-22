# from pypdf import PdfReader
# import pandas as pd

# reader = PdfReader("bank.pdf")
# if reader.is_encrypted:
#     status = reader.decrypt("7809000100028654")
#     if status > 0:
#         print(reader.pages[0].extract_text())

import openpyxl
from datetime import datetime

wb = openpyxl.load_workbook("bank_entries.xlsx", data_only=True)
sheet = wb['bank_xml_working_practice']

bank_ledger = sheet.cell(row=1, column=9).value

xml_data = """<ENVELOPE>
<HEADER>
<VERSION>1</VERSION>
<TALLYREQUEST>Import</TALLYREQUEST>
<TYPE>Data</TYPE>
<ID>Vouchers</ID>
</HEADER>
<BODY>
<DATA>
"""

def format_date(value):
    if not value:
        return ""
    if isinstance(value, str):
        return datetime.strptime(value, "%d.%m.%y").strftime("%Y%m%d")
    return value.strftime("%Y%m%d")

def payment_contra_entry_func():
    data = f"""
    <TALLYMESSAGE>
    <VOUCHER ACTION="Create">
    <VOUCHERNUMBER></VOUCHERNUMBER>
    <VOUCHERTYPENAME>Contra</VOUCHERTYPENAME>
    <DATE>{date}</DATE>
    <NARRATION>{narration}</NARRATION>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>Cash</LEDGERNAME>
    <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
    <AMOUNT>-{amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{bank_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>No</ISDEEMEDPOSITIVE>
    <AMOUNT>{amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    """
    data += "</VOUCHER></TALLYMESSAGE>"
    return data

def payment_entry_func():
    data = f"""
    <TALLYMESSAGE>
    <VOUCHER ACTION="Create">
    <VOUCHERNUMBER></VOUCHERNUMBER>
    <VOUCHERTYPENAME>Payment</VOUCHERTYPENAME>
    <DATE>{date}</DATE>
    <NARRATION>{narration}</NARRATION>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{party_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
    <AMOUNT>-{amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{bank_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>No</ISDEEMEDPOSITIVE>
    <AMOUNT>{amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    """
    data += "</VOUCHER></TALLYMESSAGE>"
    return data

def receipt_contra_entry_func():
    data = f"""
    <TALLYMESSAGE>
    <VOUCHER ACTION="Create">
    <VOUCHERNUMBER></VOUCHERNUMBER>
    <VOUCHERTYPENAME>Contra</VOUCHERTYPENAME>
    <DATE>{date}</DATE>
    <NARRATION>{narration}</NARRATION>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{bank_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
    <AMOUNT>-{amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>Cash</LEDGERNAME>
    <ISDEEMEDPOSITIVE>No</ISDEEMEDPOSITIVE>
    <AMOUNT>{amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    """
    data += "</VOUCHER></TALLYMESSAGE>"
    return data

def receipt_entry_func():
    data = f"""
    <TALLYMESSAGE>
    <VOUCHER ACTION="Create">
    <VOUCHERTYPENAME>Receipt</VOUCHERTYPENAME>
    <DATE>{date}</DATE>
    <NARRATION>{narration}</NARRATION>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{bank_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
    <AMOUNT>-{amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{party_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>No</ISDEEMEDPOSITIVE>
    <AMOUNT>{amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    """
    data += "</VOUCHER></TALLYMESSAGE>"
    return data

for row in range(2, sheet.max_row + 1):
    # Skip completely empty rows
    if not any(sheet.cell(row=row, column=col).value for col in range(1, 12)):
        continue

    print(type(sheet.cell(row=row, column=1).value), sheet.cell(row=row, column=1).value)
    date = format_date(sheet.cell(row=row, column=1).value) or ""
    
    party_ledger = sheet.cell(row=row, column=2).value or ""
    entry_type = sheet.cell(row=row, column=3).value or ""
    amount = sheet.cell(row=row, column=5).value or ""
    cheque_no = sheet.cell(row=row, column=6).value or ""
    narration_first_line = sheet.cell(row=row, column=7).value or ""

    narration_text_with_cheque = f"{narration_first_line}&#10;{sheet.cell(row=row, column=4).value}&#10;CHQ NO.: {cheque_no}"
    narration_text_without_cheque = f"{narration_first_line}&#10;{sheet.cell(row=row, column=4).value}"

    if cheque_no == "":
        narration = narration_text_without_cheque
    else:
        narration = narration_text_with_cheque
    
    if entry_type == "DR":
        if party_ledger == "Cash":
            xml_data += payment_contra_entry_func()
        else:
            xml_data += payment_entry_func()
    else:
        if party_ledger == "Cash":
            xml_data += receipt_contra_entry_func()
        else:
            xml_data += receipt_entry_func()

xml_data += """
</DATA>
</BODY>
</ENVELOPE>
"""

with open("Vch.xml", "w") as f:
    f.write(xml_data)

print("XML file created successfully")


