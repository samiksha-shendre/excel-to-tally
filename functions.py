import os
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime

#converts date to "20261019" i.e., "YYYYmmdd" format
def format_date(value):
    if not value:
        return ""
    if isinstance(value, str):
        return datetime.strptime(value, "%d/%m/%Y").strftime("%Y%m%d")
    return value.strftime("%Y%m%d")

# get 2B file of particular company
def get_2B_file(return_month, gstNo, gstr_2B_download_date):
    gstr_2B_filename = f"{return_month}_{gstNo}_GSTR2B_{gstr_2B_download_date}.xlsx"
    folder_path = Path("./CURRENT MONTH GST WORKINGS/GSTR-2B")

    if folder_path.exists():
        for file_path in folder_path.glob("*.xlsx"):
            if gstNo in file_path.name:
                return load_workbook(f"{folder_path}/{gstr_2B_filename}", data_only=True)

# get working file of particular company
def get_working_file(companyName, month, year):
    working_filename = f"GST WORKING_{companyName}_{month} {year}.xlsx"
    folder_path = Path("./CURRENT MONTH GST WORKINGS/WORKINGS")

    if folder_path.exists():
        for file_path in folder_path.glob("*.xlsx"):
            if companyName in file_path.name:
                return [load_workbook(f"{folder_path}/{working_filename}", data_only=True), folder_path]

def give_bg_color(sheet_name, min_row, max_row, min_col, max_col, bg_color):
    for row in sheet_name.iter_rows(min_row=min_row, max_row = max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.fill = bg_color

def assign_value_to_cell(sheet_name, a,b,c):
    sheet_name.cell(row=a,column=b,value=c)

def get_value_from_sheet(sheet_name,row,column):
    val = sheet_name.cell(row=row,column=column).value
    return val if val else ""

def copy_from_1_to_2(sheet_1, row_1, col_1, sheet_2, row_2, col_2):
    val = get_value_from_sheet(sheet_1,row_1,col_1)
    assign_value_to_cell(sheet_2,row_2,col_2,val)

def round_value_to_2_dec(sheet_name,start_row,end_row,start_column,end_column):
    for row in sheet_name.iter_rows(min_row=start_row, max_row=end_row, min_col=start_column, max_col=end_column):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '0.00'

def give_fixed_width(sheet_name,column,width_value):
    sheet_name.column_dimensions[column].width = width_value

def give_variable_width(sheet_name,column):
    max_len = 0

# for col in working_wb_purchases_sheet.columns:
#     max_len = 0
#     col_letter = get_column_letter(col[0].column)  

#     for cell in col:
#         if cell.value is not None:
#             # Cast the value to a string to check its character length
#             cell_len = len(str(cell.value))
#             if cell_len > max_len:
#                 max_len = cell_len
                
#     working_wb_purchases_sheet.column_dimensions[col_letter].width = max_len + 3

