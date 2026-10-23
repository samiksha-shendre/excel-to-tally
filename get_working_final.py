from openpyxl import load_workbook
from GST_NO_AND_NAMES import companies
import functions

details_wb = load_workbook("./CURRENT MONTH GST WORKINGS/DETAILS.xlsx", data_only=True)
details_sheet = details_wb["Sheet1"]
return_month = details_sheet["B8"].value
gstr_2B_download_date = details_sheet["B9"].value
gst_month = details_sheet["B1"].value
gst_year = details_sheet["B2"].value

heading_names = ["SR. NO.", "INVOICE NO.", "DATE", "PARTY NAME", "GST NO.", "NET AMOUNT", "%", "CGST", "%", "SGST", "%","IGST", "R/OFF", "TOTAL AMOUNT"]

even_width_columns = [["A", 7], ["C", 14], ["E", 22], ["G", 7], ["I", 7], ["K", 7], ["M", 7]]
odd_width_columns = ["B", "D", "F", "H", "J", "L", "N"]

for company in companies:
    compName = company["name"]
    gstNo = company["GST No."]

    wb_2B = functions.get_2B_file(return_month, gstNo, gstr_2B_download_date)
    wb_working_details = functions.get_working_file(compName, gst_month, gst_year)
    wb_working = wb_working_details[0]
    file_path = wb_working_details[1]
    wb_working_filename = f"{file_path}/GST WORKING_{compName}_{gst_month} {gst_year}.xlsx"

    sheet_2B = wb_2B["B2B"]
    sheet_wb = wb_working.create_sheet("PURCHASES")

    for row in sheet_wb['A1:A14']:
        for cell in row:
            functions.assign_value_to_cell(sheet_wb,cell.column,cell.row,heading_names[cell.row-1])
    
    functions.assign_value_to_cell(sheet_wb,2,2,gst_month)

    for row in range(7, sheet_2B.max_row + 1):
        # Skip completely empty rows
        if not any(sheet_2B.cell(row=row, column=col).value for col in range(1, 21)):
            continue
        
        #invoice_no
        functions.copy_from_1_to_2(sheet_2B, row, 3, sheet_wb, row-4, 2)
        #invoice_date
        functions.copy_from_1_to_2(sheet_2B, row, 5, sheet_wb, row-4, 3)
        #creditor_ledger
        functions.copy_from_1_to_2(sheet_2B, row, 2, sheet_wb, row-4, 4)
        #gst_no
        functions.copy_from_1_to_2(sheet_2B, row, 1, sheet_wb, row-4, 5)
        #basic_amount
        functions.copy_from_1_to_2(sheet_2B, row, 9, sheet_wb, row-4, 6)
        #cgst_amount
        functions.copy_from_1_to_2(sheet_2B, row, 11, sheet_wb, row-4, 8)
        #sgst_amount
        functions.copy_from_1_to_2(sheet_2B, row, 12, sheet_wb, row-4, 10)
        #igst_amount
        functions.copy_from_1_to_2(sheet_2B, row, 10, sheet_wb, row-4, 12)
        #invoice_value
        functions.copy_from_1_to_2(sheet_2B, row, 6, sheet_wb, row-4, 14)
        #rounding_off_value
        rounding_off_amt = float(functions.get_value_from_sheet(sheet_2B,row,6)) - float(functions.get_value_from_sheet(sheet_2B,row,9)) - float(functions.get_value_from_sheet(sheet_2B,row,10) or 0) - float(functions.get_value_from_sheet(sheet_2B,row,11) or 0) - float(functions.get_value_from_sheet(sheet_2B,row,12) or 0)

        functions.assign_value_to_cell(sheet_wb, row-4,13,rounding_off_amt) if rounding_off_amt else ""
        
    functions.round_value_to_2_dec(sheet_wb,3,sheet_wb.max_row,6,14)

    for col, wid in even_width_columns:
        functions.give_fixed_width(sheet_wb, col, wid)
    

    wb_working.save(wb_working_filename)
