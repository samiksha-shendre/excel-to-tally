import openpyxl
from datetime import datetime

wb = openpyxl.load_workbook("GSTR-2B.xlsx", data_only=True)
month_sheet = wb['Read me']
sheet = wb['B2B']

narration_first_line = "BEING MATERIAL PURCHASED FROM PARTY"
# rounding_off_ledger = "" 
tax_month = month_sheet.cell(row=9, column=3).value
tax_period_month_number = datetime.strptime(tax_month, "%d/%m/%Y").month - 1
tax_period_year_number = datetime.strptime(tax_month, "%d/%m/%Y").year
base_date = format_date(f"01/{tax_period_month_number}/{tax_period_year_number}")

def format_date(value):
    if not value:
        return ""
    if isinstance(value, str):
        return datetime.strptime(value, "%d/%m/%Y").strftime("%Y%m%d")
    return value.strftime("%Y%m%d")

def igst_purchase_entry_func(invoice_no, invoice_date, creditor_ledger, basic_amount, igst_amount, invoice_value, purchase_ledger, igst_ledger):
    # write logic to compare dates and write narration accordingly.
    formatted_invoice_date = format_date(invoice_date)

    if base_date > formatted_invoice_date:
        narration_text = f"{narration_first_line}&#10;BILL DATE : {invoice_date}"
        date = base_date
    else:
        narration_text = f"{narration_first_line}"
        date = invoice_date

    data = f"""
    <TALLYMESSAGE>
    <VOUCHER ACTION="Create">
    <VOUCHERNUMBER></VOUCHERNUMBER>
    <VOUCHERTYPENAME>Purchase</VOUCHERTYPENAME>
    <DATE>{date}</DATE>
    <REFERENCEDATE>{invoice_date}</REFERENCEDATE>
    <REFERENCE>{invoice_no}</REFERENCE>
    <NARRATION>{narration_text}</NARRATION>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{creditor_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>No</ISDEEMEDPOSITIVE>
    <AMOUNT>{invoice_value}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>

    ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{purchase_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
    <AMOUNT>-{basic_amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
   
    ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{igst_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
    <AMOUNT>-{igst_amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    """

    #rounding_off logic
    basic_igst_total = basic_amount + igst_amount
    if basic_igst_total != invoice_value:
        if basic_igst_total < invoice_value:
            rounding_off_amt = invoice_value - basic_igst_total
            data += f"""
            ALLLEDGERENTRIES.LIST>
            <LEDGERNAME>{rounding_off_ledger}</LEDGERNAME>
            <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
            <AMOUNT>-{rounding_off_amt}</AMOUNT>
            </ALLLEDGERENTRIES.LIST>
            """
        else:
            rounding_off_amt = basic_igst_total - invoice_value
            data += f"""
            ALLLEDGERENTRIES.LIST>
            <LEDGERNAME>{rounding_off_ledger}</LEDGERNAME>
            <ISDEEMEDPOSITIVE>No</ISDEEMEDPOSITIVE>
            <AMOUNT>{rounding_off_amt}</AMOUNT>
            </ALLLEDGERENTRIES.LIST>
            """ 
    data += "</VOUCHER></TALLYMESSAGE>"
    return data

def cgst_sgst_purchase_entry_func(invoice_no, invoice_date, creditor_ledger, basic_amount, cgst_amount, sgst_amount, invoice_value, purchase_ledger, cgst_ledger, sgst_ledger):
    formatted_invoice_date = format_date(invoice_date)

    if base_date > formatted_invoice_date:
        narration_text = f"{narration_first_line}&#10;BILL DATE : {invoice_date}"
        date = base_date
    else:
        narration_text = f"{narration_first_line}"
        date = invoice_date

    data = f"""
    <TALLYMESSAGE>
    <VOUCHER ACTION="Create">
    <VOUCHERNUMBER></VOUCHERNUMBER>
    <VOUCHERTYPENAME>Purchase</VOUCHERTYPENAME>
    <DATE>{date}</DATE>
    <REFERENCEDATE>{invoice_date}</REFERENCEDATE>
    <REFERENCE>{invoice_no}</REFERENCE>
    <NARRATION>{narration_text}</NARRATION>

    <ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{creditor_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>No</ISDEEMEDPOSITIVE>
    <AMOUNT>{invoice_value}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>

    ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{purchase_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
    <AMOUNT>-{basic_amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
   
    ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{cgst_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
    <AMOUNT>-{cgst_amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>

    ALLLEDGERENTRIES.LIST>
    <LEDGERNAME>{sgst_ledger}</LEDGERNAME>
    <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
    <AMOUNT>-{sgst_amount}</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    """

    #rounding_off logic
    basic_cgst_sgst_total = basic_amount + cgst_amount + sgst_amount
    if basic_cgst_sgst_total != invoice_value:
        if basic_cgst_sgst_total < invoice_value:
            rounding_off_amt = invoice_value - basic_cgst_sgst_total
            data += f"""
            ALLLEDGERENTRIES.LIST>
            <LEDGERNAME>{rounding_off_ledger}</LEDGERNAME>
            <ISDEEMEDPOSITIVE>Yes</ISDEEMEDPOSITIVE>
            <AMOUNT>-{rounding_off_amt}</AMOUNT>
            </ALLLEDGERENTRIES.LIST>
            """
        else:
            rounding_off_amt = basic_cgst_sgst_total - invoice_value
            data += f"""
            ALLLEDGERENTRIES.LIST>
            <LEDGERNAME>{rounding_off_ledger}</LEDGERNAME>
            <ISDEEMEDPOSITIVE>No</ISDEEMEDPOSITIVE>
            <AMOUNT>{rounding_off_amt}</AMOUNT>
            </ALLLEDGERENTRIES.LIST>
            """ 
    data += "</VOUCHER></TALLYMESSAGE>"
    return data

xml_data = """<ENVELOPE>
<HEADER>
<VERSION>1</VERSION>
<TALLYREQUEST>Import</TALLYREQUEST>
<TYPE>Data</TYPE>
<ID>Vouchers</ID>
</HEADER>
<BODY>
<DATA>
"""

for row in range(7, sheet.max_row + 1):
    # Skip completely empty rows
    if not any(sheet.cell(row=row, column=col).value for col in range(1, 21)):
        continue

    invoice_no = sheet.cell(row=row, column=3).value or ""
    invoice_date = format_date(sheet.cell(row=row, column=5).value) or ""
    creditor_ledger = sheet.cell(row=row, column=2).value or ""
    # gst_no = sheet.cell(row=row, column=1).value or ""
    basic_amount = sheet.cell(row=row, column=9).value or ""
    cgst_amount = sheet.cell(row=row, column=11).value or ""
    sgst_amount = sheet.cell(row=row, column=12).value or ""
    igst_amount = sheet.cell(row=row, column=10).value or ""
    invoice_value = sheet.cell(row=row, column=6).value or ""
    # purchase_ledger = ""
    # cgst_ledger = ""  
    # sgst_ledger = "" 
    # igst_ledger = ""

    if igst_amount == 0:
        cgst_sgst_purchase_entry_func(invoice_no, invoice_date, creditor_ledger, basic_amount, cgst_amount, sgst_amount, invoice_value, purchase_ledger, cgst_ledger, sgst_ledger)
    else:
        igst_purchase_entry_func(invoice_no, invoice_date, creditor_ledger, basic_amount, igst_amount, invoice_value, purchase_ledger, igst_ledger)


#     rate = (cgst_amount + sgst_amount) / basic_amount if not igst_amount else igst_amount / basic_amount

#     match rate:
#         case 0.18:
#             return purchase_ledger = "PURCHASE @ 18%"
#         case 0.05:
#             return purchase_ledger = "PURCHASE @ 5%"
#         case 0.12:
#             return purchase_ledger = "PURCHASE @ 12%"
#         case 0.28:
#             return purchase_ledger = "PURCHASE @ 28%"

#     print(rate)
    
#     party_ledger = sheet.cell(row=row, column=2).value or ""
#     entry_type = sheet.cell(row=row, column=3).value or ""
#     amount = sheet.cell(row=row, column=5).value or ""
#     cheque_no = sheet.cell(row=row, column=6).value or ""
#     narration_first_line = sheet.cell(row=row, column=7).value or ""

#     narration_text_with_cheque = f"{narration_first_line}&#10;{sheet.cell(row=row, column=4).value}&#10;CHQ NO.: {cheque_no}"
#     narration_text_without_cheque = f"{narration_first_line}&#10;{sheet.cell(row=row, column=4).value}"

#     if cheque_no == "":
#         narration = narration_text_without_cheque
#     else:
#         narration = narration_text_with_cheque
    
#     if entry_type == "DR":
#         if party_ledger == "Cash":
#             xml_data += payment_contra_entry_func()
#         else:
#             xml_data += payment_entry_func()
#     else:
#         if party_ledger == "Cash":
#             xml_data += receipt_contra_entry_func()
#         else:
#             xml_data += receipt_entry_func()

# xml_data += """
# </DATA>
# </BODY>
# </ENVELOPE>
# """

# with open("Vch.xml", "w") as f:
#     f.write(xml_data)

# print("XML file created successfully")

